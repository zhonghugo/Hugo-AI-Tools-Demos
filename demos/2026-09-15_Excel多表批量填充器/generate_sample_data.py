# -*- coding: utf-8 -*-
"""
示例数据生成器
================
运行本脚本会生成 demo 所需的全部示例文件：
  1. sample_data/合同清单.xlsx   —— 源表（8条合同数据，模拟"从合同/清单里扒数据"）
  2. sample_data/对内台账.xlsx   —— 目标表1（内部管理台账，含公式列+合计行）
  3. sample_data/对外报表.xlsx   —— 目标表2（对外报送，含序号公式）
  4. sample_data/合规检查表.xlsx —— 目标表3（合规检查，含"检查结论"待填列）
  5. sample_data/付款计划表.xlsx —— 目标表4（付款计划，含比例×金额联动公式）

用法：
    python3 generate_sample_data.py
生成后，运行 python3 batch_filler.py 即可看到批量填充效果。
"""
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------- 通用样式 ----------
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def _style_body(ws, r1, r2, ncols):
    for r in range(r1, r2 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER


def create_source_table():
    """源表：合同清单（模拟从合同台账/清单里扒出来的原始数据）"""
    wb = Workbook()
    ws = wb.active
    ws.title = "合同明细"
    headers = ["合同编号", "项目名称", "供应商", "合同金额", "签订日期",
               "付款方式", "负责人", "备注"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))

    rows = [
        ["HT-2026-001", "A厂区消防改造工程", "中安消防工程有限公司", 1860000, datetime.date(2026, 3, 12), "分三期", "王强", "已归档"],
        ["HT-2026-002", "B楼智能化弱电项目", "华讯智能科技", 932000, datetime.date(2026, 4, 2), "预付30%", "李敏", ""],
        ["HT-2026-003", "C园区绿化养护服务", "绿洲园林", 458000, datetime.date(2026, 4, 18), "按季度", "赵磊", "服务期1年"],
        ["HT-2026-004", "D数据中心UPS电源采购", "恒动电力设备", 2740000, datetime.date(2026, 5, 6), "预付50%", "王强", ""],
        ["HT-2026-005", "E办公楼中央空调维保", "舒适冷暖工程", 316000, datetime.date(2026, 5, 20), "按年度", "孙悦", "含备件"],
        ["HT-2026-006", "F园区安防监控改造", "视界安防", 689000, datetime.date(2026, 6, 15), "分两期", "李敏", "新增点位32个"],
        ["HT-2026-007", "G厂区地面翻新工程", "新基建设", 1280000, datetime.date(2026, 7, 1), "预付40%", "赵磊", ""],
        ["HT-2026-008", "H办公楼电梯维保服务", "迅达电梯服务", 235000, datetime.date(2026, 7, 22), "按年度", "孙悦", "含年检"],
    ]
    for row in rows:
        ws.append(row)
    _style_body(ws, 2, 1 + len(rows), len(headers))

    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=5).number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=4).number_format = "#,##0"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 18

    wb.save("sample_data/合同清单.xlsx")
    print("✓ 已生成源表 sample_data/合同清单.xlsx（8条合同）")


def _make_target_sheet(wb, sheet_name, headers, ncols, nrows=12):
    """建一张带表头的目标表，返回 (ws, 数据起始行)"""
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    _style_header(ws, 1, ncols)
    # 预置 nrows 行边框，模拟"已经排版好的固定模板"
    _style_body(ws, 2, 1 + nrows, ncols)
    return ws, 2


def create_target_1():
    """对内台账：合同金额由填充写入，含税金额=金额×1.13（公式保留），底部合计=SUM"""
    wb = Workbook()
    ws, start = _make_target_sheet(wb, "台账", ["合同编号", "项目名称", "供应商", "合同金额",
                                               "含税金额(13%)", "签订日期", "负责人"], 7)
    # 含税金额列（E列）预置公式，模拟"模板里写好的公式，不希望被覆盖"
    for r in range(start, start + 12):
        ws.cell(row=r, column=5).value = f"=D{r}*1.13"
        ws.cell(row=r, column=5).number_format = "#,##0.00"
    # 合计行（第15行）
    tr = start + 13
    ws.cell(row=tr, column=2).value = "合计"
    ws.cell(row=tr, column=4).value = f"=SUM(D{start}:D{tr - 1})"
    ws.cell(row=tr, column=5).value = f"=SUM(E{start}:E{tr - 1})"
    for c in range(1, 8):
        cell = ws.cell(row=tr, column=c)
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.border = BORDER
    ws.cell(row=tr, column=4).number_format = "#,##0"
    ws.cell(row=tr, column=5).number_format = "#,##0.00"

    for col, w in zip("ABCDEFG", [14, 26, 22, 14, 16, 12, 10]):
        ws.column_dimensions[col].width = w
    wb.save("sample_data/对内台账.xlsx")
    print("✓ 已生成目标表1 sample_data/对内台账.xlsx（含税公式+合计行）")


def create_target_2():
    """对外报表：序号列用公式 ROW()-1 自动编号，其余列由填充写入"""
    wb = Workbook()
    ws, start = _make_target_sheet(wb, "对外报表", ["序号", "合同编号", "项目名称", "合同金额",
                                                  "签订日期", "备注"], 6)
    for r in range(start, start + 12):
        ws.cell(row=r, column=1).value = f"=ROW()-{start - 1}"
    for col, w in zip("ABCDEF", [6, 14, 26, 14, 12, 18]):
        ws.column_dimensions[col].width = w
    wb.save("sample_data/对外报表.xlsx")
    print("✓ 已生成目标表2 sample_data/对外报表.xlsx（序号自动编号公式）")


def create_target_3():
    """合规检查表：检查结论列默认填"待查"，付款方式、负责人一起带过去"""
    wb = Workbook()
    ws, start = _make_target_sheet(wb, "合规检查", ["合同编号", "项目名称", "供应商", "合同金额",
                                                 "付款方式", "检查结论", "负责人"], 7)
    for r in range(start, start + 12):
        ws.cell(row=r, column=6).value = "待查"
    for col, w in zip("ABCDEFG", [14, 26, 22, 14, 12, 12, 10]):
        ws.column_dimensions[col].width = w
    wb.save("sample_data/合规检查表.xlsx")
    print("✓ 已生成目标表3 sample_data/合规检查表.xlsx（检查结论默认'待查'）")


def create_target_4():
    """付款计划表：首付款比例写入，首付款金额=合同金额×比例（公式联动）"""
    wb = Workbook()
    ws, start = _make_target_sheet(wb, "付款计划", ["合同编号", "项目名称", "供应商", "合同金额",
                                                 "首付款比例", "首付款金额"], 6)
    for r in range(start, start + 12):
        ws.cell(row=r, column=6).value = f"=ROUND(D{r}*E{r},2)"
        ws.cell(row=r, column=6).number_format = "#,##0.00"
        ws.cell(row=r, column=5).number_format = "0%"
    for col, w in zip("ABCDEF", [14, 26, 22, 14, 12, 14]):
        ws.column_dimensions[col].width = w
    wb.save("sample_data/付款计划表.xlsx")
    print("✓ 已生成目标表4 sample_data/付款计划表.xlsx（首付款金额公式联动）")


if __name__ == "__main__":
    create_source_table()
    create_target_1()
    create_target_2()
    create_target_3()
    create_target_4()
    print("\n示例数据全部生成完毕。现在运行：python3 batch_filler.py")
