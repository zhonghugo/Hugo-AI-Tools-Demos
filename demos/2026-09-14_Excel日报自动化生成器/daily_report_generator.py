#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel日报自动化生成器
=====================
解决问题：每天重复做Excel日报（清洗数据→计算指标→做图→输出），步骤固定但繁琐。
使用方法：配置一次config.json，以后每天只需放入新数据，运行本脚本即可自动生成完整日报。

运行命令：
    python3 daily_report_generator.py
    python3 daily_report_generator.py --config my_config.json
    python3 daily_report_generator.py --input 今日数据.xlsx --output 今日日报.xlsx
"""

import json
import argparse
import os
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows


# ============================================================
# 第一部分：数据清洗
# ============================================================
def clean_data(df, cleaning_config):
    """
    根据配置清洗数据
    - 去重
    - 空值填充
    - 日期格式统一
    - 去除前后空格
    """
    print("  [1/4] 正在清洗数据...")

    # 1. 去除字符串列的前后空格
    for col in df.select_dtypes(include=['object']).columns:
        df[col] = df[col].astype(str).str.strip()

    # 2. 去重
    if cleaning_config.get('drop_duplicates', False):
        before = len(df)
        df = df.drop_duplicates().copy()
        print(f"    去重：{before} → {len(df)} 行（移除 {before - len(df)} 条重复）")

    # 3. 空值填充
    fillna = cleaning_config.get('fillna', {})
    for col, value in fillna.items():
        if col in df.columns:
            filled = df[col].isna().sum()
            df[col] = df[col].fillna(value)
            if filled > 0:
                print(f"    空值填充：列「{col}」填充 {filled} 个空值为 {value}")

    # 4. 日期格式统一
    date_col = cleaning_config.get('date_column', '')
    if date_col and date_col in df.columns:
        df[date_col] = pd.to_datetime(df[date_col], errors='coerce')
        print(f"    日期格式化：列「{date_col}」已统一为日期格式")

    print(f"    清洗完成：共 {len(df)} 行，{len(df.columns)} 列")
    return df


# ============================================================
# 第二部分：指标计算
# ============================================================
def calculate_metrics(df, metrics_config, group_by=None):
    """
    根据配置计算指标
    支持：sum（求和）、count（计数）、mean（均值）、max（最大值）、min（最小值）
    如果配置了group_by，则按维度分组计算
    """
    print("  [2/4] 正在计算指标...")

    results = {}
    formula_map = {
        'sum': 'sum',
        'count': 'count',
        'mean': 'mean',
        'max': 'max',
        'min': 'min'
    }

    # 总体指标
    overall = {}
    for metric in metrics_config:
        name = metric['name']
        formula = metric['formula']
        column = metric['column']

        if column not in df.columns:
            print(f"    ⚠️  列「{column}」不存在，跳过指标「{name}」")
            continue

        if formula == 'sum':
            value = df[column].sum()
        elif formula == 'count':
            value = df[column].count()
        elif formula == 'mean':
            value = df[column].mean()
        elif formula == 'max':
            value = df[column].max()
        elif formula == 'min':
            value = df[column].min()
        else:
            value = None

        overall[name] = value
        print(f"    {name}: {value:,.2f}" if isinstance(value, (int, float)) else f"    {name}: {value}")

    results['overall'] = overall

    # 分组指标
    if group_by and group_by in df.columns:
        grouped = df.groupby(group_by)
        group_results = {}
        for metric in metrics_config:
            name = metric['name']
            formula = formula_map.get(metric['formula'], 'sum')
            column = metric['column']
            if column in df.columns:
                group_results[name] = grouped[column].agg(formula).to_dict()
        results['grouped'] = group_results
        print(f"    已按「{group_by}」分组计算指标")

    return results


# ============================================================
# 第三部分：生成Excel日报
# ============================================================
def generate_report(df, metrics, config, output_path):
    """
    生成包含多个Sheet的Excel日报：
    - Sheet1: 日报概览（指标汇总+生成时间）
    - Sheet2: 清洗后数据
    - Sheet3: 分组指标（含图表）
    """
    print("  [3/4] 正在生成Excel日报...")

    group_by = config.get('group_by', '')
    charts_config = config.get('charts', [])

    # 用pandas先写基础数据
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        # Sheet1: 日报概览
        overview_data = {'指标': list(metrics['overall'].keys()), '数值': list(metrics['overall'].values())}
        overview_df = pd.DataFrame(overview_data)
        overview_df.to_excel(writer, sheet_name='日报概览', index=False, startrow=2)

        # Sheet2: 清洗后数据
        df.to_excel(writer, sheet_name='清洗后数据', index=False)

        # Sheet3: 分组指标
        if 'grouped' in metrics and group_by:
            group_df = pd.DataFrame(metrics['grouped'])
            group_df.index.name = group_by
            group_df.to_excel(writer, sheet_name='分组指标')

    # 用openpyxl美化格式 + 添加图表
    wb = load_workbook(output_path)

    # --- 美化日报概览 ---
    ws_overview = wb['日报概览']
    # 标题
    ws_overview['A1'] = f'📊 每日数据日报 - {datetime.now().strftime("%Y-%m-%d")}'
    ws_overview['A1'].font = Font(size=16, bold=True, color='FFFFFF')
    ws_overview['A1'].fill = PatternFill(start_color='2B579A', end_color='2B579A', fill_type='solid')
    ws_overview.merge_cells('A1:B1')
    ws_overview['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws_overview.row_dimensions[1].height = 30

    # 表头样式
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(bold=True, color='FFFFFF')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for cell in ws_overview[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    # 数据行样式
    for row in ws_overview.iter_rows(min_row=4, max_row=ws_overview.max_row, max_col=2):
        for cell in row:
            cell.border = thin_border
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'

    ws_overview.column_dimensions['A'].width = 20
    ws_overview.column_dimensions['B'].width = 20

    # --- 美化清洗后数据 ---
    ws_data = wb['清洗后数据']
    for cell in ws_data[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center')
    ws_data.freeze_panes = 'A2'
    for col in ws_data.columns:
        max_length = max(len(str(cell.value or '')) for cell in col)
        ws_data.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)

    # --- 添加图表到分组指标 ---
    if '分组指标' in wb.sheetnames and charts_config:
        ws_chart = wb['分组指标']

        for i, chart_cfg in enumerate(charts_config):
            chart_type = chart_cfg.get('type', 'bar')
            title = chart_cfg.get('title', '图表')

            if chart_type == 'bar':
                chart = BarChart()
                chart.type = 'col'
            elif chart_type == 'line':
                chart = LineChart()
            elif chart_type == 'pie':
                chart = PieChart()
            else:
                chart = BarChart()

            chart.title = title
            chart.style = 10
            chart.y_axis.title = '数值'
            chart.x_axis.title = chart_cfg.get('x', '类别')

            # 数据范围（第2列开始为指标数据，第1列为分组维度）
            data = Reference(ws_chart, min_col=2, min_row=1,
                             max_col=ws_chart.max_column, max_row=ws_chart.max_row)
            cats = Reference(ws_chart, min_col=1, min_row=2, max_row=ws_chart.max_row)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.width = 18
            chart.height = 10

            # 图表放在数据右侧
            chart_position = f'D{2 + i * 18}'
            ws_chart.add_chart(chart, chart_position)

        print(f"    已添加 {len(charts_config)} 个图表")

    wb.save(output_path)
    print(f"    日报已保存：{output_path}")
    return output_path


# ============================================================
# 第四部分：主函数
# ============================================================
def main():
    parser = argparse.ArgumentParser(description='Excel日报自动化生成器')
    parser.add_argument('--config', default='config.json', help='配置文件路径（默认：config.json）')
    parser.add_argument('--input', help='输入数据文件路径（覆盖配置文件中的设置）')
    parser.add_argument('--output', help='输出日报文件路径（覆盖配置文件中的设置）')
    args = parser.parse_args()

    # 读取配置
    config_path = Path(args.config)
    if not config_path.exists():
        print(f"❌ 配置文件不存在：{config_path}")
        print("   请先创建config.json，或运行 create_sample_data.py 生成示例配置和数据")
        return

    with open(config_path, 'r', encoding='utf-8') as f:
        config = json.load(f)

    # 命令行参数覆盖
    if args.input:
        config['input_file'] = args.input
    if args.output:
        config['output_file'] = args.output

    input_file = config['input_file']
    output_file = config['output_file']

    print("=" * 60)
    print("  Excel日报自动化生成器")
    print("=" * 60)
    print(f"  输入文件：{input_file}")
    print(f"  输出文件：{output_file}")
    print()

    # 检查输入文件
    if not Path(input_file).exists():
        print(f"❌ 输入文件不存在：{input_file}")
        print("   请运行 create_sample_data.py 生成示例数据，或修改config.json中的input_file")
        return

    # 读取数据
    sheet_name = config.get('sheet_name', 0)
    try:
        df = pd.read_excel(input_file, sheet_name=sheet_name)
        print(f"✅ 数据读取成功：{len(df)} 行 × {len(df.columns)} 列")
        print(f"   列名：{', '.join(df.columns.tolist())}")
    except Exception as e:
        print(f"❌ 数据读取失败：{e}")
        return

    print()

    # 步骤1：清洗数据
    cleaning_config = config.get('cleaning', {})
    df = clean_data(df, cleaning_config)

    # 步骤2：计算指标
    metrics_config = config.get('metrics', [])
    group_by = config.get('group_by', '')
    metrics = calculate_metrics(df, metrics_config, group_by)

    # 步骤3：生成日报
    output_path = generate_report(df, metrics, config, output_file)

    # 步骤4：完成
    print()
    print("  [4/4] 完成！")
    print("=" * 60)
    print(f"  ✅ 日报已生成：{os.path.abspath(output_path)}")
    print(f"  📊 包含Sheet：日报概览 / 清洗后数据 / 分组指标（含图表）")
    print(f"  ⏱️  生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print("=" * 60)
    print()
    print("  💡 后续使用：每天只需替换 input_file 为新数据，重新运行即可")
    print("  🔧 如需调整指标或图表，编辑 config.json 即可")


if __name__ == '__main__':
    main()
